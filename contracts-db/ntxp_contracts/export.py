"""Export the Master Contracts database to an Excel workbook.

One workbook, organized around the contract: the log, coefficient tiers,
contacts, every job grouped under the contract it was done under (unlinked jobs
last), owner-entity contact profiles, and the sources behind each record. Sheets
print on Letter portrait, fit to one page wide, header rows repeated.

Needs openpyxl (``pip install 'ntxp-contracts[xlsx]'``); imported lazily so the
core package stays dependency-free.
"""
from __future__ import annotations

import datetime as _dt
import json
from pathlib import Path

from . import branding
from .db.repository import Repository

_MONEY = '"$"#,##0'


def export_xlsx(repo: Repository, path: str | Path, open_questions: list[str] | None = None) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as exc:  # pragma: no cover
        raise SystemExit("openpyxl is required: pip install 'ntxp-contracts[xlsx]'") from exc

    today = _dt.date.today()
    # Same brand the screens use — resolved from the brand skill at export time.
    colors = branding.resolve_brand().colors
    navy = colors["primary"].lstrip("#")
    tint = colors["light"].lstrip("#") if colors["light"].upper() != "#FFFFFF" else "E8EDF4"
    alert, warn, link_c = (colors[k].lstrip("#") for k in ("danger", "warning", "accent2"))
    conn = repo.conn
    # Live vehicles first, then those in motion, then history; soonest expiry first.
    rank = {"awarded": 0, "pending-execution": 1, "tentative": 2, "unknown": 3, "expired": 4, "lost": 5}
    contracts = sorted(repo.list_contracts(), key=lambda c: (
        rank.get(c.get("status"), 3), c.get("expiration_date") or "9999", c["contract_title"]))
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=navy)
    group_fill = PatternFill("solid", fgColor=tint)

    def sheet(title: str, heading: str, columns: list[tuple[str, int]], first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        n = len(columns)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        ws.cell(1, 1, heading).font = Font(bold=True, size=14, color=navy)
        ws.cell(1, 1).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
        ws.cell(2, 1, f"NTXP LLC · generated {today:%B %d, %Y}").alignment = Alignment(horizontal="center")
        for i, (label, width) in enumerate(columns, 1):
            cell = ws.cell(3, i, label)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A4"
        ws.print_title_rows = "3:3"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddFooter.center.text = f"{title} — page &P of &N"
        return ws

    def put(ws, row, values, money_cols=(), bold=False, fill=None):
        for i, v in enumerate(values, 1):
            cell = ws.cell(row, i, v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if i in money_cols:
                cell.number_format = _MONEY
            if bold:
                cell.font = Font(bold=True)
            if fill:
                cell.fill = fill

    def link(ws, row, col, url, text):
        if url:
            cell = ws.cell(row, col, text)
            cell.hyperlink = url
            cell.font = Font(color=link_c, underline="single")

    # 1 — Contract log --------------------------------------------------------
    ws = sheet("Contract Log", "MASTER CONTRACTS LOG", [
        ("Contract #", 14), ("Contract title", 34), ("Type", 16), ("Owner entity", 24),
        ("Award status", 13), ("Signed?", 10), ("Award date", 11), ("Expires", 11),
        ("Days left", 8), ("Coefficient", 10), ("Co-op fee", 10), ("Pricebook", 14),
        ("Allowable scope", 30), ("Jobs", 6), ("Job value", 13), ("Invoiced", 13),
        ("Signed copy", 12)], first=True)
    r = 4
    for c in contracts:
        prof = repo.get_profile(c["contract_id"])
        t = prof["totals"]
        days = None
        if c.get("expiration_date"):
            try:
                days = (_dt.date.fromisoformat(c["expiration_date"]) - today).days
            except ValueError:
                pass
        put(ws, r, [c.get("contract_no"), c["contract_title"], c.get("contract_type"),
                    c.get("owner_entity"), c.get("status"),
                    "Yes" if c["is_executed"] else "NO — unexecuted",
                    c.get("award_date"), c.get("expiration_date"), days,
                    c.get("coefficient_multiplier"), c.get("cooperative_fee"),
                    c.get("pricebook"), c.get("allowable_scope"), t["job_count"],
                    t["total_contract_value"] or None, t["total_sales"] or None, None],
            money_cols=(15, 16))
        if not c["is_executed"]:
            ws.cell(r, 6).font = Font(bold=True, color=alert)
        if days is not None and days <= 90:
            ws.cell(r, 9).font = Font(bold=True, color=alert if days < 0 else warn)
        link(ws, r, 17, c.get("pdf_url"), "Open")
        r += 1
    ws.auto_filter.ref = f"A3:Q{max(r - 1, 3)}"

    # 2 — Coefficients --------------------------------------------------------
    ws = sheet("Coefficients", "COEFFICIENTS BY CONTRACT", [
        ("Contract #", 14), ("Contract title", 40), ("Tier", 32), ("Coefficient", 12)])
    r = 4
    for c in contracts:
        tiers = repo.list_coefficients(c["contract_id"]) or (
            [{"label": "Primary", "value": c["coefficient_multiplier"]}]
            if c.get("coefficient_multiplier") is not None else [])
        for t in tiers:
            put(ws, r, [c.get("contract_no"), c["contract_title"], t["label"], t["value"]])
            r += 1

    # 3 — Contacts ------------------------------------------------------------
    ws = sheet("Contacts", "CONTRACT CONTACTS", [
        ("Contract #", 14), ("Owner entity", 24), ("Name", 22), ("Title", 22),
        ("Role", 15), ("Email", 28), ("Phone", 15), ("Source", 9)])
    r = 4
    for c in contracts:
        for ct in repo.list_contacts(c["contract_id"], c.get("owner_id")):
            put(ws, r, [c.get("contract_no"), ct.get("org") or c.get("owner_entity"),
                        ct["name"], ct.get("title"), ct.get("role"), ct.get("email"),
                        ct.get("phone"), None])
            link(ws, r, 8, ct.get("source_url"), "Open")
            r += 1

    # 4 — Jobs grouped by contract -------------------------------------------
    cols = [("Job #", 9), ("Job", 34), ("Customer", 24), ("Status", 11), ("Created", 11),
            ("Closed", 11), ("Contract value", 13), ("Invoiced", 13), ("Paid", 13),
            ("Link", 8), ("Why linked", 34)]
    ws = sheet("Jobs by Contract", "JOBS BY CONTRACT", cols)
    r = 4

    def job_rows(title, jobs):
        nonlocal r
        put(ws, r, [title] + [None] * (len(cols) - 1), bold=True, fill=group_fill)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        r += 1
        for j in jobs:
            put(ws, r, [j.get("job_number"), j["name"], j.get("customer"), j.get("status"),
                        j.get("created_on"), j.get("closed_on"), j.get("contract_value"),
                        j.get("sales_amount"), j.get("paid_amount"), j.get("link_confidence"),
                        j.get("link_reason")], money_cols=(7, 8, 9))
            r += 1
        put(ws, r, [None, f"{len(jobs)} job(s)", None, None, None, "Subtotal",
                    sum(j.get("contract_value") or 0 for j in jobs) or None,
                    sum(j.get("sales_amount") or 0 for j in jobs) or None,
                    sum(j.get("paid_amount") or 0 for j in jobs) or None, None, None],
            money_cols=(7, 8, 9), bold=True)
        r += 2

    for c in contracts:
        jobs = repo.list_jobs(c["contract_id"])
        label = f"{c.get('contract_no') or '(no #)'} — {c['contract_title']} ({c.get('owner_entity') or ''})"
        job_rows(label, jobs)
    unlinked = repo.list_unlinked_jobs()
    if unlinked:
        job_rows("NOT LINKED TO A CONTRACT VEHICLE (direct / subcontract / private work)", unlinked)

    # 5 — Owner entities ------------------------------------------------------
    ws = sheet("Owner Entities", "OWNER ENTITY CONTACT PROFILES", [
        ("Owner entity", 28), ("Website", 26), ("Main", 15), ("Customer service", 15),
        ("Accounting", 15), ("Email", 26), ("Address", 34)])
    r = 4
    for o in conn.execute("SELECT * FROM owner_entities ORDER BY name").fetchall():
        put(ws, r, [o["name"], o["website"], o["main_phone"], o["customer_service_phone"],
                    o["accounting_phone"], o["email"], o["address"]])
        link(ws, r, 2, o["website"], o["website"])
        r += 1

    # 6 — Sources -------------------------------------------------------------
    ws = sheet("Sources", "WHERE EACH RECORD CAME FROM", [
        ("Contract #", 14), ("Contract title", 30), ("Evidence", 52), ("Link", 9)])
    r = 4
    for c in contracts:
        for s in repo.list_sources(c["contract_id"]):
            put(ws, r, [c.get("contract_no"), c["contract_title"], s["what"], None])
            link(ws, r, 4, s.get("url"), "Open")
            r += 1

    # 7 — Open questions ------------------------------------------------------
    if open_questions:
        ws = sheet("Open Questions", "OPEN QUESTIONS TO RESOLVE", [("#", 5), ("Question", 100)])
        for i, q in enumerate(open_questions, 1):
            put(ws, 3 + i, [i, q])

    out = Path(path)
    wb.save(out)
    return out


def export_json(repo: Repository) -> dict:
    """Everything, as plain JSON — the shape other tools poll or re-import."""
    out = []
    for c in repo.list_contracts():
        out.append(repo.get_profile(c["contract_id"]))
    return {"generated": _dt.date.today().isoformat(), "contracts": out,
            "unlinked_jobs": repo.list_unlinked_jobs(), "stats": repo.stats()}


def dumps(obj) -> str:
    return json.dumps(obj, indent=2, default=str)
